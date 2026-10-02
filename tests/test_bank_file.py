"""Tests for bank file generation and the audit report."""

from __future__ import annotations

import io
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from pavt.bank_file import AUDIT_COLUMNS, BANK_COLUMNS, build_audit_report, build_bank_file
from pavt.models import InvoiceDocument, MatchStatus, PaymentRequest, SourceKind
from pavt.verification import verify_payment

FIXED_TIME = datetime(2026, 9, 24, 10, 30, 0)


def make_result(request_no, payee, account, requested, invoice_amount, bank_code="004"):
    request = PaymentRequest(
        request_no=request_no,
        payee_name=payee,
        bank_code=bank_code,
        bank_account=account,
        particulars="Service Fee",
        amount=Decimal(requested),
    )
    invoice = (
        None
        if invoice_amount is None
        else InvoiceDocument(
            source=Path(f"{request_no}.pdf"),
            total_amount=Decimal(invoice_amount),
            payee_name=payee,
            bank_account=account,
            source_kind=SourceKind.OCR_LOCAL,
        )
    )
    return verify_payment(request, invoice, tolerance=Decimal("0.01"))


@pytest.fixture(scope="module")
def sample_results():
    return [
        make_result("P000001", "ABC Medical Centre Limited", "123456789", "15000.00", "15000.00"),
        make_result("P000002", "XYZ Office Supplies Limited", "456789123", "5800.00", "5800.00"),
        make_result("P000003", "Healthy Imaging Centre Limited", "987654321", "125000.00", "120000.00", bank_code="024"),
    ]


def load_sheet(content: bytes, title: str | None = None):
    workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    return workbook[title] if title else workbook.worksheets[0]


class TestBankFile:
    def test_columns_match_the_required_format(self, sample_results):
        sheet = load_sheet(build_bank_file(sample_results, generated_at=FIXED_TIME).content)
        header = [cell.value for cell in sheet[1]]
        assert header == [title for title, _ in BANK_COLUMNS]
        assert header[0] == "Payment Request No"
        assert header[1] == "Payee Name"
        assert header[2] == "Bank Code and Account Number"
        assert header[3] == "Payment Amount (HKD)"

    def test_only_verified_payments_are_included(self, sample_results):
        bank = build_bank_file(sample_results, generated_at=FIXED_TIME)
        assert bank.row_count == 2
        sheet = load_sheet(bank.content)
        payees = [sheet.cell(row=row, column=2).value for row in (2, 3)]
        assert payees == ["ABC Medical Centre Limited", "XYZ Office Supplies Limited"]
        # The mismatched supplier must be absent entirely.
        all_values = [
            cell.value for row in sheet.iter_rows() for cell in row if cell.value is not None
        ]
        assert "Healthy Imaging Centre Limited" not in all_values

    def test_bank_code_and_account_are_combined(self, sample_results):
        sheet = load_sheet(build_bank_file(sample_results, generated_at=FIXED_TIME).content)
        assert sheet.cell(row=2, column=3).value == "004-123456789"

    def test_payee_details_follow_the_summary(self):
        """Payee/account come from the summary even when the invoice differs."""
        request = PaymentRequest(
            request_no="P1",
            payee_name="Summary Name Limited",
            bank_code="004",
            bank_account="111222333",
            particulars="Fee",
            amount=Decimal("100.00"),
        )
        invoice = InvoiceDocument(
            source=Path("inv.pdf"),
            total_amount=Decimal("100.00"),
            payee_name="Different Invoice Name",
            bank_account="999999999",
            source_kind=SourceKind.TEXT_LAYER,
        )
        result = verify_payment(request, invoice)
        sheet = load_sheet(build_bank_file([result], generated_at=FIXED_TIME).content)
        assert sheet.cell(row=2, column=2).value == "Summary Name Limited"
        assert sheet.cell(row=2, column=3).value == "004-111222333"

    def test_amount_uses_the_crosschecked_value(self, sample_results):
        sheet = load_sheet(build_bank_file(sample_results, generated_at=FIXED_TIME).content)
        assert sheet.cell(row=2, column=4).value == 15000
        assert sheet.cell(row=2, column=4).number_format == "#,##0.00"

    def test_total_row_and_totals(self, sample_results):
        bank = build_bank_file(sample_results, generated_at=FIXED_TIME)
        assert bank.total_amount == Decimal("20800.00")
        sheet = load_sheet(bank.content)
        assert sheet.cell(row=4, column=3).value == "TOTAL"
        assert sheet.cell(row=4, column=4).value == 20800

    def test_exclusions_are_summarised_by_reason(self, sample_results):
        bank = build_bank_file(sample_results, generated_at=FIXED_TIME)
        assert bank.excluded_count == 1
        assert bank.excluded_reasons == {"Amount mismatch": 1}
        summary = load_sheet(bank.content, "Verification Summary")
        values = [cell.value for row in summary.iter_rows() for cell in row]
        assert "Amount mismatch" in values

    def test_empty_batch_still_produces_a_valid_workbook(self):
        bank = build_bank_file([], batch_reference="EMPTY", generated_at=FIXED_TIME)
        assert bank.row_count == 0
        sheet = load_sheet(bank.content)
        assert sheet.cell(row=2, column=3).value == "TOTAL"
        assert sheet.cell(row=2, column=4).value == 0

    def test_filename_uses_the_batch_reference(self, sample_results):
        bank = build_bank_file(sample_results, batch_reference="BATCH-7", generated_at=FIXED_TIME)
        assert bank.filename == "BATCH-7_bank_payment_file.xlsx"

    def test_write_to_disk(self, sample_results, tmp_path):
        bank = build_bank_file(sample_results, generated_at=FIXED_TIME)
        written = bank.write_to(tmp_path / "nested" / bank.filename)
        assert written.exists()
        assert written.stat().st_size > 0


class TestAuditReport:
    def test_reports_every_row(self, sample_results):
        audit = build_audit_report(sample_results)
        assert audit.row_count == 3
        sheet = load_sheet(audit.content)
        assert [cell.value for cell in sheet[1]] == [title for title, _ in AUDIT_COLUMNS]

    def test_holds_are_labelled(self, sample_results):
        sheet = load_sheet(build_audit_report(sample_results).content)
        statuses = {sheet.cell(row=row, column=1).value: sheet.cell(row=row, column=8).value
                    for row in (2, 3, 4)}
        assert statuses["P000001"] == "Matched"
        assert statuses["P000003"] == "Amount mismatch"

    def test_difference_is_shown(self, sample_results):
        sheet = load_sheet(build_audit_report(sample_results).content)
        assert sheet.cell(row=4, column=7).value == 5000

    def test_no_invoice_rows_are_reported(self):
        result = make_result("P1", "ABC Ltd", "123", "100.00", None)
        sheet = load_sheet(build_audit_report([result]).content)
        assert sheet.cell(row=2, column=8).value == "No invoice"
        # Empty cells read back as None from a saved workbook.
        assert sheet.cell(row=2, column=5).value is None
