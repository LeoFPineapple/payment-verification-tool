"""Bank upload file generation (Excel) and the audit report.

The bank file contains **only** payments that passed verification -- a mismatched
or unpaired request never reaches it.  Payee name, bank code and account number
are copied from the Payment Request Summary (the file of record); the amount is
the crosschecked figure.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import datetime
from decimal import Decimal
from pathlib import Path

from .models import MatchStatus, VerificationResult
from .money import format_amount, quantize

BANK_COLUMNS = (
    ("Payment Request No", 18),
    ("Payee Name", 34),
    ("Bank Code and Account Number", 30),
    ("Payment Amount (HKD)", 22),
)

AUDIT_COLUMNS = (
    ("Payment Request No", 18),
    ("Payee Name", 34),
    ("Bank Code and Account Number", 30),
    ("Requested Amount (HKD)", 20),
    ("Invoice File", 24),
    ("Invoice Amount (HKD)", 20),
    ("Difference (HKD)", 16),
    ("Status", 18),
    ("Read Via", 16),
    ("Match Confidence", 16),
    ("Remarks", 60),
)

_STATUS_LABEL = {
    MatchStatus.MATCHED: "Matched",
    MatchStatus.AMOUNT_MISMATCH: "Amount mismatch",
    MatchStatus.NO_INVOICE: "No invoice",
    MatchStatus.UNREADABLE: "Unreadable invoice",
}


@dataclass(slots=True)
class BankFile:
    """A generated bank file (or audit report) held in memory."""

    filename: str
    content: bytes
    row_count: int = 0
    total_amount: Decimal = Decimal("0.00")
    excluded_count: int = 0
    excluded_reasons: dict[str, int] = field(default_factory=dict)

    def write_to(self, path: str | Path) -> Path:
        """Persist the workbook to disk."""
        destination = Path(path)
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_bytes(self.content)
        return destination


def _style_header(worksheet, columns) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(bold=True, color="FFFFFF")
    alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for index, (title, width) in enumerate(columns, start=1):
        cell = worksheet.cell(row=1, column=index, value=title)
        cell.fill = fill
        cell.font = font
        cell.alignment = alignment
        worksheet.column_dimensions[get_column_letter(index)].width = width
    worksheet.freeze_panes = "A2"


def build_bank_file(
    results: list[VerificationResult],
    *,
    batch_reference: str | None = None,
    generated_at: datetime | None = None,
) -> BankFile:
    """Build the bank payment file from verification results.

    Args:
        results: Verification output for the batch.
        batch_reference: Optional label used in the file name and summary sheet.
        generated_at: Timestamp override, mainly so tests are deterministic.

    Returns:
        A :class:`BankFile` whose ``content`` is a ready-to-download ``.xlsx``.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    included = [r for r in results if r.included_in_bank_file]
    excluded = [r for r in results if not r.included_in_bank_file]
    stamp = generated_at or datetime.now()
    reference = batch_reference or f"PAYMENT-{stamp:%Y%m%d-%H%M%S}"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Bank Payment File"
    _style_header(sheet, BANK_COLUMNS)

    total = Decimal("0.00")
    for offset, result in enumerate(included):
        row = offset + 2
        request = result.request
        # Amount per the crosscheck result: the invoice total that was verified.
        amount = quantize(result.invoice_amount if result.invoice_amount is not None else request.amount)
        total += amount
        sheet.cell(row=row, column=1, value=request.request_no)
        sheet.cell(row=row, column=2, value=request.payee_name)
        sheet.cell(row=row, column=3, value=request.bank_code_and_account)
        amount_cell = sheet.cell(row=row, column=4, value=float(amount))
        amount_cell.number_format = "#,##0.00"
        amount_cell.alignment = Alignment(horizontal="right")

    total_row = len(included) + 2
    label = sheet.cell(row=total_row, column=3, value="TOTAL")
    label.font = Font(bold=True)
    total_cell = sheet.cell(row=total_row, column=4, value=float(total))
    total_cell.font = Font(bold=True)
    total_cell.number_format = "#,##0.00"
    total_cell.alignment = Alignment(horizontal="right")
    for column in range(1, len(BANK_COLUMNS) + 1):
        sheet.cell(row=total_row, column=column).fill = PatternFill("solid", fgColor="DDEBF7")

    # --- summary sheet: what went in, what was held back and why ---------------
    reasons: dict[str, int] = {}
    for result in excluded:
        reasons[_STATUS_LABEL[result.status]] = reasons.get(_STATUS_LABEL[result.status], 0) + 1

    summary = workbook.create_sheet("Verification Summary")
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 46
    summary_rows = [
        ("Batch reference", reference),
        ("Generated at", stamp.strftime("%Y-%m-%d %H:%M:%S")),
        ("Payment requests in summary", len(results)),
        ("Payments included in bank file", len(included)),
        ("Payments excluded", len(excluded)),
        ("Total amount payable (HKD)", format_amount(total)),
    ]
    for offset, (label_text, value) in enumerate(summary_rows, start=1):
        key_cell = summary.cell(row=offset, column=1, value=label_text)
        key_cell.font = Font(bold=True)
        summary.cell(row=offset, column=2, value=value)
    offset = len(summary_rows) + 2
    heading = summary.cell(row=offset, column=1, value="Exclusions by reason")
    heading.font = Font(bold=True)
    for reason, count in sorted(reasons.items()):
        offset += 1
        summary.cell(row=offset, column=1, value=reason)
        summary.cell(row=offset, column=2, value=count)
    if not reasons:
        summary.cell(row=offset + 1, column=1, value="None - all requests verified")

    buffer = io.BytesIO()
    workbook.save(buffer)

    return BankFile(
        filename=f"{reference}_bank_payment_file.xlsx",
        content=buffer.getvalue(),
        row_count=len(included),
        total_amount=total,
        excluded_count=len(excluded),
        excluded_reasons=reasons,
    )


def build_audit_report(results: list[VerificationResult]) -> BankFile:
    """Build an audit workbook listing every request and its verdict."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Verification Audit"
    _style_header(sheet, AUDIT_COLUMNS)

    ok_fill = PatternFill("solid", fgColor="E2EFDA")
    hold_fill = PatternFill("solid", fgColor="FCE4E4")

    for offset, result in enumerate(results):
        row = offset + 2
        request = result.request
        values = [
            request.request_no,
            request.payee_name,
            request.bank_code_and_account,
            float(quantize(request.amount)),
            result.invoice.filename if result.invoice else "",
            float(quantize(result.invoice_amount)) if result.invoice_amount is not None else "",
            float(quantize(result.difference)) if result.difference is not None else "",
            _STATUS_LABEL[result.status],
            result.invoice.source_kind.value if result.invoice else "",
            round(result.match_confidence, 2),
            (result.note + (f" Match basis: {result.match_reason}." if result.match_reason else "")).strip(),
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=column, value=value)
            if column in (4, 6, 7):
                cell.number_format = "#,##0.00"
                cell.alignment = Alignment(horizontal="right")
            if column in (8, 10):
                cell.alignment = Alignment(horizontal="center")
        fill = ok_fill if result.included_in_bank_file else hold_fill
        sheet.cell(row=row, column=8).fill = fill
        sheet.cell(row=row, column=8).font = Font(bold=True)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return BankFile(
        filename="Verification_Audit_Report.xlsx",
        content=buffer.getvalue(),
        row_count=len(results),
    )
